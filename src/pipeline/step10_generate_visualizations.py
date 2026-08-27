"""Phase 2 - Step 10: Visualization exports (interactive HTML + Excel).

Builds the following visualizations and writes them to two Power BI /
analyst-friendly formats:

1. Training set health-state distribution
2. Confusion matrix of the training set (BP Neural Network)
3. Test set health-state distribution
4. Confusion matrix of the test set (BP Neural Network)
5. Sample points vs Mahalanobis Distance
6. Sample points vs Health Index
7. Mahalanobis Distance vs Health Index
8. Sample points vs Health Score (AHI)

Outputs:
    data/exports/visualizations.html  (interactive Plotly dashboard)
    data/exports/visualizations.xlsx  (data tables + native Excel charts)
"""

from __future__ import annotations

import numpy as np
import pandas as pd
import plotly.graph_objects as go
from openpyxl.chart import BarChart, Reference, ScatterChart, Series
from openpyxl.formatting.rule import ColorScaleRule
from plotly.subplots import make_subplots
from sklearn.metrics import confusion_matrix

from src.pipeline.common import VISUALIZATIONS_EXCEL_PATH, VISUALIZATIONS_HTML_PATH
from src.utils.logger import get_logger

logger = get_logger(__name__)


def _label_distribution(labels: np.ndarray, class_order: list[str]) -> pd.Series:
    counts = pd.Series(labels).value_counts()
    return counts.reindex(class_order, fill_value=0)


def _confusion_df(
    y_true: np.ndarray, y_pred: np.ndarray, class_order: list[str]
) -> pd.DataFrame:
    matrix = confusion_matrix(y_true, y_pred, labels=class_order)
    return pd.DataFrame(matrix, index=class_order, columns=class_order)


def build_visualization_data(
    split_info: dict, ahi_scores_df: pd.DataFrame
) -> dict[str, pd.DataFrame]:
    """Assemble the tabular datasets backing every requested visualization.

    Args:
        split_info: ``split_info`` dict returned by Step 4
            (``labels``, ``y_train``, ``y_train_pred``, ``y_test``, ``y_test_pred``).
        ahi_scores_df: Output of Step 6 (``asset_id``, ``timestamp``,
            ``mahalanobis_distance``, ``health_index``, ``ahi_score``).

    Returns:
        A dictionary of named DataFrames, one per visualization/data table.
    """
    class_order = sorted(split_info["labels"])

    train_dist = _label_distribution(split_info["y_train"], class_order)
    test_dist = _label_distribution(split_info["y_test"], class_order)

    train_confusion = _confusion_df(
        split_info["y_train"], split_info["y_train_pred"], class_order
    )
    test_confusion = _confusion_df(
        split_info["y_test"], split_info["y_test_pred"], class_order
    )

    sample_scores = ahi_scores_df.copy().reset_index(drop=True)
    sample_scores.insert(0, "sample_index", range(1, len(sample_scores) + 1))
    sample_scores["sample_label"] = (
        sample_scores["sample_index"].astype(str) + " - " + sample_scores["asset_id"]
    )

    return {
        "train_distribution": train_dist.rename("count").rename_axis("health_status").reset_index(),
        "test_distribution": test_dist.rename("count").rename_axis("health_status").reset_index(),
        "train_confusion_matrix": train_confusion,
        "test_confusion_matrix": test_confusion,
        "sample_scores": sample_scores,
    }


def _build_html_dashboard(data: dict[str, pd.DataFrame]) -> go.Figure:
    """Build a single interactive Plotly dashboard with all 8 charts."""
    fig = make_subplots(
        rows=4,
        cols=2,
        subplot_titles=(
            "1. Training Set Health-State Distribution",
            "2. Training Set Confusion Matrix",
            "3. Test Set Health-State Distribution",
            "4. Test Set Confusion Matrix",
            "5. Sample Points vs Mahalanobis Distance",
            "6. Sample Points vs Health Index",
            "7. Mahalanobis Distance vs Health Index",
            "8. Sample Points vs Health Score (AHI)",
        ),
        specs=[
            [{"type": "bar"}, {"type": "heatmap"}],
            [{"type": "bar"}, {"type": "heatmap"}],
            [{"type": "xy"}, {"type": "xy"}],
            [{"type": "xy"}, {"type": "xy"}],
        ],
        vertical_spacing=0.08,
        horizontal_spacing=0.12,
    )

    train_dist = data["train_distribution"]
    fig.add_trace(
        go.Bar(x=train_dist["health_status"], y=train_dist["count"], name="Train"),
        row=1,
        col=1,
    )

    train_cm = data["train_confusion_matrix"]
    fig.add_trace(
        go.Heatmap(
            z=train_cm.values,
            x=list(train_cm.columns),
            y=list(train_cm.index),
            colorscale="Blues",
            text=train_cm.values,
            texttemplate="%{text}",
            showscale=False,
            name="Train Confusion Matrix",
        ),
        row=1,
        col=2,
    )

    test_dist = data["test_distribution"]
    fig.add_trace(
        go.Bar(x=test_dist["health_status"], y=test_dist["count"], name="Test"),
        row=2,
        col=1,
    )

    test_cm = data["test_confusion_matrix"]
    fig.add_trace(
        go.Heatmap(
            z=test_cm.values,
            x=list(test_cm.columns),
            y=list(test_cm.index),
            colorscale="Oranges",
            text=test_cm.values,
            texttemplate="%{text}",
            showscale=False,
            name="Test Confusion Matrix",
        ),
        row=2,
        col=2,
    )

    scores = data["sample_scores"]
    fig.add_trace(
        go.Scatter(
            x=scores["sample_label"],
            y=scores["mahalanobis_distance"],
            mode="markers+lines",
            name="Mahalanobis Distance",
        ),
        row=3,
        col=1,
    )
    fig.add_trace(
        go.Scatter(
            x=scores["sample_label"],
            y=scores["health_index"],
            mode="markers+lines",
            name="Health Index",
        ),
        row=3,
        col=2,
    )
    fig.add_trace(
        go.Scatter(
            x=scores["mahalanobis_distance"],
            y=scores["health_index"],
            mode="markers",
            text=scores["sample_label"],
            name="MD vs HI",
        ),
        row=4,
        col=1,
    )
    fig.add_trace(
        go.Scatter(
            x=scores["sample_label"],
            y=scores["ahi_score"],
            mode="markers+lines",
            name="Health Score (AHI)",
        ),
        row=4,
        col=2,
    )

    fig.update_xaxes(title_text="Health Status", row=1, col=1)
    fig.update_yaxes(title_text="Count", row=1, col=1)
    fig.update_xaxes(title_text="Health Status", row=2, col=1)
    fig.update_yaxes(title_text="Count", row=2, col=1)
    fig.update_xaxes(title_text="Sample Point", row=3, col=1, tickangle=45)
    fig.update_yaxes(title_text="Mahalanobis Distance", row=3, col=1)
    fig.update_xaxes(title_text="Sample Point", row=3, col=2, tickangle=45)
    fig.update_yaxes(title_text="Health Index", row=3, col=2)
    fig.update_xaxes(title_text="Mahalanobis Distance", row=4, col=1)
    fig.update_yaxes(title_text="Health Index", row=4, col=1)
    fig.update_xaxes(title_text="Sample Point", row=4, col=2, tickangle=45)
    fig.update_yaxes(title_text="Health Score (AHI)", row=4, col=2)

    fig.update_layout(
        height=1600,
        width=1400,
        showlegend=False,
        title_text="Asset Health Index - Model & Scoring Visualizations",
    )
    return fig


def _write_excel_workbook(data: dict[str, pd.DataFrame], path) -> None:
    """Write data tables plus native Excel charts to a single workbook."""
    with pd.ExcelWriter(path, engine="openpyxl") as writer:
        data["train_distribution"].to_excel(writer, sheet_name="Train Distribution", index=False)
        data["test_distribution"].to_excel(writer, sheet_name="Test Distribution", index=False)
        data["train_confusion_matrix"].to_excel(writer, sheet_name="Train Confusion Matrix")
        data["test_confusion_matrix"].to_excel(writer, sheet_name="Test Confusion Matrix")
        data["sample_scores"].to_excel(writer, sheet_name="Sample Scores", index=False)

        workbook = writer.book

        # --- Bar charts: train/test health-state distribution ---
        for sheet_name in ("Train Distribution", "Test Distribution"):
            ws = workbook[sheet_name]
            n_rows = ws.max_row
            chart = BarChart()
            chart.title = sheet_name
            chart.x_axis.title = "Health Status"
            chart.y_axis.title = "Count"
            data_ref = Reference(ws, min_col=2, min_row=1, max_row=n_rows)
            categories_ref = Reference(ws, min_col=1, min_row=2, max_row=n_rows)
            chart.add_data(data_ref, titles_from_data=True)
            chart.set_categories(categories_ref)
            ws.add_chart(chart, "E2")

        # --- Conditional-formatted heatmaps for the confusion matrices ---
        for sheet_name, n_classes in (
            ("Train Confusion Matrix", len(data["train_confusion_matrix"])),
            ("Test Confusion Matrix", len(data["test_confusion_matrix"])),
        ):
            ws = workbook[sheet_name]
            last_col_letter = ws.cell(row=1, column=n_classes + 1).column_letter
            cell_range = f"B2:{last_col_letter}{n_classes + 1}"
            rule = ColorScaleRule(
                start_type="min", start_color="FFFFFF",
                end_type="max", end_color="4472C4",
            )
            ws.conditional_formatting.add(cell_range, rule)

        # --- Scatter charts: sample points / MD / HI / AHI ---
        ws = workbook["Sample Scores"]
        n_rows = ws.max_row
        header = [cell.value for cell in ws[1]]
        col = {name: header.index(name) + 1 for name in header}

        def _scatter_chart(title: str, y_col: str, x_col: str, y_title: str, x_title: str) -> ScatterChart:
            chart = ScatterChart()
            chart.title = title
            chart.x_axis.title = x_title
            chart.y_axis.title = y_title
            x_values = Reference(ws, min_col=col[x_col], min_row=2, max_row=n_rows)
            y_values = Reference(ws, min_col=col[y_col], min_row=1, max_row=n_rows)
            series = Series(y_values, x_values, title_from_data=True)
            series.marker.symbol = "circle"
            series.graphicalProperties.line.noFill = True
            chart.series.append(series)
            return chart

        ws.add_chart(
            _scatter_chart(
                "5. Sample Points vs Mahalanobis Distance",
                "mahalanobis_distance",
                "sample_index",
                "Mahalanobis Distance",
                "Sample Point",
            ),
            "K2",
        )
        ws.add_chart(
            _scatter_chart(
                "6. Sample Points vs Health Index",
                "health_index",
                "sample_index",
                "Health Index",
                "Sample Point",
            ),
            "K20",
        )
        ws.add_chart(
            _scatter_chart(
                "7. Mahalanobis Distance vs Health Index",
                "health_index",
                "mahalanobis_distance",
                "Health Index",
                "Mahalanobis Distance",
            ),
            "K38",
        )
        ws.add_chart(
            _scatter_chart(
                "8. Sample Points vs Health Score (AHI)",
                "ahi_score",
                "sample_index",
                "Health Score (AHI)",
                "Sample Point",
            ),
            "K56",
        )


def run_visualization_export(split_info: dict, ahi_scores_df: pd.DataFrame) -> dict:
    """Build and export all 8 visualizations as interactive HTML + Excel.

    Args:
        split_info: ``split_info`` dict returned by Step 4.
        ahi_scores_df: Output of Step 6.

    Returns:
        Mapping of ``{"html": path, "excel": path}``.
    """
    data = build_visualization_data(split_info, ahi_scores_df)

    VISUALIZATIONS_HTML_PATH.parent.mkdir(parents=True, exist_ok=True)
    dashboard = _build_html_dashboard(data)
    dashboard.write_html(VISUALIZATIONS_HTML_PATH, include_plotlyjs="cdn")
    logger.info("Wrote interactive HTML dashboard to '%s'.", VISUALIZATIONS_HTML_PATH)

    VISUALIZATIONS_EXCEL_PATH.parent.mkdir(parents=True, exist_ok=True)
    _write_excel_workbook(data, VISUALIZATIONS_EXCEL_PATH)
    logger.info("Wrote Excel visualization workbook to '%s'.", VISUALIZATIONS_EXCEL_PATH)

    return {"html": VISUALIZATIONS_HTML_PATH, "excel": VISUALIZATIONS_EXCEL_PATH}


if __name__ == "__main__":
    from src.pipeline.step2_preprocess_data import run_preprocessing
    from src.pipeline.step4_train_bp_network import run_bp_training
    from src.pipeline.step5_score_mahalanobis import run_mahalanobis_scoring
    from src.pipeline.step6_calculate_health_index import run_health_index_calculation

    _clean_df, _preprocessor = run_preprocessing()
    _, _metrics, _split_info = run_bp_training(_clean_df, _preprocessor)
    _md_scores_df = run_mahalanobis_scoring()
    _ahi_scores_df = run_health_index_calculation(_md_scores_df)
    run_visualization_export(_split_info, _ahi_scores_df)
